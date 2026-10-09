"""
Test script to verify document parsing (PDF, DOCX, TXT), chunking, and FAISS indexing locally.
"""
import os
import docx
from reportlab.pdfgen import canvas
from rag_engine import load_document, split_documents, get_embedding_model, create_vector_store, retrieve_relevant_chunks

def create_sample_files():
    os.makedirs("sample_docs", exist_ok=True)
    
    # 1. Create TXT file
    txt_path = os.path.join("sample_docs", "artificial_intelligence.txt")
    with open(txt_path, "w", encoding="utf-8") as f:
        f.write(
            "Artificial Intelligence (AI) Overview:\n"
            "AI is a branch of computer science dedicated to developing systems capable of performing tasks "
            "that typically require human intelligence. Key subfields include machine learning, deep learning, "
            "natural language processing (NLP), and computer vision. In 2026, Retrieval-Augmented Generation (RAG) "
            "has become a mainstream approach for grounding LLMs on private enterprise documents without retraining."
        )
    print(f"Created {txt_path}")

    # 2. Create DOCX file
    docx_path = os.path.join("sample_docs", "project_guidelines.docx")
    doc = docx.Document()
    doc.add_heading("Company Project Guidelines 2026", level=1)
    doc.add_paragraph("All engineering projects must maintain strict unit test coverage of at least 85%.")
    doc.add_paragraph("Deployment to production happens every Tuesday at 10:00 AM UTC.")
    doc.save(docx_path)
    print(f"Created {docx_path}")

    # 3. Create PDF file
    pdf_path = os.path.join("sample_docs", "cloud_infrastructure.pdf")
    c = canvas.Canvas(pdf_path)
    c.drawString(100, 750, "Cloud Infrastructure Architecture Report")
    c.drawString(100, 720, "Our primary cloud region is AWS us-east-1 with automated failover to us-west-2.")
    c.drawString(100, 690, "The vector database cluster uses FAISS for low latency in-memory vector search.")
    c.save()
    print(f"Created {pdf_path}")

    return [txt_path, docx_path, pdf_path]

def run_tests():
    files = create_sample_files()
    all_docs = []
    
    print("\n--- Testing Document Loading ---")
    for file_path in files:
        docs = load_document(file_path)
        print(f"Loaded '{file_path}': {len(docs)} page(s)/section(s)")
        all_docs.extend(docs)
        
    print(f"\nTotal documents loaded: {len(all_docs)}")
    
    print("\n--- Testing Document Chunking ---")
    chunks = split_documents(all_docs, chunk_size=300, chunk_overlap=50)
    print(f"Total chunks created: {len(chunks)}")
    for i, chunk in enumerate(chunks):
        print(f" Chunk {i+1} [{chunk.metadata['source']} p.{chunk.metadata['page']}]: {chunk.page_content[:60]}...")
        
    print("\n--- Testing Hugging Face Embeddings & FAISS ---")
    embedding_model = get_embedding_model()
    vector_store = create_vector_store(chunks, embedding_model)
    print("FAISS vector store created successfully!")
    
    print("\n--- Testing Similarity Search ---")
    queries = [
        "When does deployment to production happen?",
        "What is the primary cloud region?",
        "What is RAG used for in AI?"
    ]
    for q in queries:
        print(f"\nQuery: '{q}'")
        results = retrieve_relevant_chunks(vector_store, q, top_k=2)
        for doc, score in results:
            print(f"  -> Match [{doc.metadata['source']}]: {doc.page_content} (score: {score:.4f})")
    print("\n[SUCCESS] ALL LOCAL PIPELINE TESTS PASSED SUCCESSFULLY!")

if __name__ == "__main__":
    run_tests()
